
import csv
import json
import os
import re
import sqlite3
import tempfile

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from generators import iterar_registros


FORMATOS = {
    "sql": ".sql",
    "db": ".db",
    "json": ".json",
    "csv": ".csv",
    "xlsx": ".xlsx",
}


def pasta_downloads():
    # localiza a pasta Downloads do usuario

    if os.name == "nt":
        try:
            import ctypes
            from ctypes import wintypes
            from uuid import UUID

            class GUID(ctypes.Structure):
                _fields_ = [
                    ("Data1", wintypes.DWORD),
                    ("Data2", wintypes.WORD),
                    ("Data3", wintypes.WORD),
                    ("Data4", ctypes.c_ubyte * 8),
                ]

            identificador = UUID(
                "374DE290-123F-4565-9164-39C4925E467B"
            )

            guid = GUID.from_buffer_copy(
                identificador.bytes_le
            )

            caminho = ctypes.c_wchar_p()

            funcao = ctypes.windll.shell32.SHGetKnownFolderPath
            funcao.argtypes = [
                ctypes.POINTER(GUID),
                wintypes.DWORD,
                wintypes.HANDLE,
                ctypes.POINTER(ctypes.c_wchar_p),
            ]
            funcao.restype = ctypes.c_long

            resultado = funcao(
                ctypes.byref(guid),
                0,
                None,
                ctypes.byref(caminho),
            )

            if resultado == 0:
                try:
                    return Path(caminho.value)
                finally:
                    ctypes.windll.ole32.CoTaskMemFree(
                        ctypes.cast(caminho, ctypes.c_void_p)
                    )

        except (OSError, AttributeError, ValueError):
            pass

    downloads = Path.home() / "Downloads"
    downloads.mkdir(parents=True, exist_ok=True)

    return downloads


def nome_arquivo(config, formato):
    tabela = config["table"]

    nome = re.sub(r"[^A-Za-z0-9_-]", "_", tabela)

    if not nome:
        nome = "dados"

    data = datetime.now().strftime("%Y%m%d_%H%M%S")

    return f"{nome}_{data}.{formato}"


def converter_json(valor):
    if isinstance(valor, Decimal):
        return float(valor)

    if isinstance(valor, (datetime,)):
        return valor.isoformat()

    raise TypeError(
        f"Tipo não serializável: {type(valor).__name__}"
    )


def exportar_json(config, quantidade, caminho):
    with open(caminho, "w", encoding="utf-8") as arquivo:
        arquivo.write("[\n")

        for indice, registro in enumerate(
            iterar_registros(config, quantidade)
        ):
            if indice:
                arquivo.write(",\n")

            arquivo.write(
                "    " + json.dumps(
                    registro,
                    ensure_ascii=False,
                    default=converter_json,
                )
            )

        arquivo.write("\n]\n")


def exportar_csv(config, quantidade, caminho):
    nomes = [
        coluna["name"]
        for coluna in config["columns"]
    ]

    with open(
        caminho,
        "w",
        encoding="utf-8-sig",
        newline=""
    ) as arquivo:

        escritor = csv.DictWriter(
            arquivo,
            fieldnames=nomes,
            delimiter=";",
        )

        escritor.writeheader()

        for registro in iterar_registros(config, quantidade):
            escritor.writerow(registro)


def identificador_sql(nome):
    return '"' + nome.replace('"', '""') + '"'


def valor_sql(valor):
    if valor is None:
        return "NULL"

    if isinstance(valor, bool):
        return "1" if valor else "0"

    if isinstance(valor, (int, float, Decimal)):
        return str(valor)

    texto = str(valor).replace("'", "''")

    return f"'{texto}'"


def criar_tabela_sql(config):
    tabela = identificador_sql(config["table"])

    colunas = []

    for coluna in config["columns"]:
        nome = identificador_sql(coluna["name"])
        tipo = coluna["type"]

        colunas.append(f"    {nome} {tipo}")

    estrutura = ",\n".join(colunas)

    return (
        f"CREATE TABLE IF NOT EXISTS {tabela} (\n"
        f"{estrutura}\n"
        f");\n"
    )


def exportar_sql(config, quantidade, caminho):
    tabela = identificador_sql(config["table"])

    nomes = [
        coluna["name"]
        for coluna in config["columns"]
    ]

    colunas_sql = ", ".join(
        identificador_sql(nome)
        for nome in nomes
    )

    with open(caminho, "w", encoding="utf-8") as arquivo:
        arquivo.write("-- Populate Database\n")
        arquivo.write("-- Dados fictícios para testes\n\n")

        arquivo.write(criar_tabela_sql(config))
        arquivo.write("\n")

        for registro in iterar_registros(config, quantidade):
            valores = ", ".join(
                valor_sql(registro[nome])
                for nome in nomes
            )

            arquivo.write(
                f"INSERT INTO {tabela} ({colunas_sql}) "
                f"VALUES ({valores});\n"
            )


def exportar_sqlite(config, quantidade, caminho):
    conexao = sqlite3.connect(caminho)

    try:
        conexao.execute(criar_tabela_sql(config))

        nomes = [
            coluna["name"]
            for coluna in config["columns"]
        ]

        tabela = identificador_sql(config["table"])

        colunas_sql = ", ".join(
            identificador_sql(nome)
            for nome in nomes
        )

        placeholders = ", ".join(
            "?" for _ in nomes
        )

        comando = (
            f"INSERT INTO {tabela} ({colunas_sql}) "
            f"VALUES ({placeholders})"
        )

        lote = []

        for registro in iterar_registros(config, quantidade):
            lote.append(
                tuple(registro[nome] for nome in nomes)
            )

            if len(lote) >= 1000:
                conexao.executemany(comando, lote)
                lote.clear()

        if lote:
            conexao.executemany(comando, lote)

        conexao.commit()

    except Exception:
        conexao.rollback()
        raise

    finally:
        conexao.close()


def exportar_excel(config, quantidade, caminho):
    # limite fisico do excel: 1.048.576 linhas, incluindo a linha do cabeçalho.
    if quantidade > 1_048_575:
        raise ValueError(
            "O Excel permite no máximo 1.048.575 registros "
            "por planilha, além do cabeçalho. "
            "Utilize CSV, JSON, SQL ou SQLite."
        )

    workbook = Workbook(write_only=True)
    planilha = workbook.create_sheet(
        title=config["table"][:31]
    )

    nomes = [
        coluna["name"]
        for coluna in config["columns"]
    ]

    planilha.append(nomes)

    for registro in iterar_registros(config, quantidade):
        planilha.append([
            registro[nome]
            for nome in nomes
        ])

    workbook.save(caminho)


def exportar(config, quantidade, formato):
    # exporta para Downloads e retorna o caminho criado
    # nunca sobrescreve um arquivo existente

    if formato not in FORMATOS:
        raise ValueError("Formato de exportação inválido.")

    quantidade = int(quantidade)

    if quantidade < 1:
        raise ValueError(
            "A quantidade precisa ser maior que zero."
        )

    if not config.get("columns"):
        raise ValueError(
            "Adicione pelo menos uma coluna."
        )

    pasta = pasta_downloads()
    pasta.mkdir(parents=True, exist_ok=True)

    nome = nome_arquivo(config, formato)

    # reserva um nome exclusivo para evitar sobrescrever
    base = Path(nome).stem
    extensao = FORMATOS[formato]

    contador = 1

    while True:
        nome_atual = (
            nome if contador == 1
            else f"{base}_{contador}{extensao}"
        )

        destino = pasta / nome_atual

        try:
            descritor = os.open(
                destino,
                os.O_CREAT | os.O_EXCL | os.O_WRONLY,
                0o600,
            )
            os.close(descritor)
            break

        except FileExistsError:
            contador += 1

    funcoes = {
        "json": exportar_json,
        "csv": exportar_csv,
        "sql": exportar_sql,
        "db": exportar_sqlite,
        "xlsx": exportar_excel,
    }

    try:
        funcoes[formato](
            config,
            quantidade,
            destino,
        )

    except Exception:
        destino.unlink(missing_ok=True)
        raise

    return destino
