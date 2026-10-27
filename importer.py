
import csv
import json
import re
import sqlite3

from pathlib import Path
from openpyxl import load_workbook
from catalog import (GERADORES, TIPOS_SQL, detectar_gerador, tipo_do_gerador, normalizar_coluna)

IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


def normalizar_nome(nome):
    nome = str(nome).strip()
    nome = re.sub(r"\W+", "_", nome, flags=re.ASCII)
    nome = nome.strip("_")

    if not nome:
        nome = "coluna"

    if nome[0].isdigit():
        nome = "_" + nome

    return nome

def nome_tabela_arquivo(caminho):
    # obtem o nome da tabela a partir do arquivo, removendo timestamps

    nome = Path(caminho).stem

    # remove o YYYYMMDD_HHMMSS do final
    nome = re.sub(
        r"_\d{8}_\d{6}$",
        "",
        nome
    )

    # remove sufixos de arquivos duplicados
    nome = re.sub(
        r"_\d{8}_\d{6}_\d+$",
        "",
        nome
    )

    return normalizar_nome(nome)


def criar_colunas(nomes):
    colunas = []
    usados = set()

    for indice, nome in enumerate(nomes, 1):
        nome = normalizar_nome(nome)

        if not nome:
            nome = f"coluna_{indice}"

        original = nome
        contador = 2

        while nome.lower() in usados:
            nome = f"{original}_{contador}"
            contador += 1

        usados.add(nome.lower())

        gerador = detectar_gerador(nome)

        colunas.append({
            "name": nome,
            "type": tipo_do_gerador(gerador),
            "generator": gerador,
            "options": {},
        })

    return colunas


def importar_csv(caminho):
    with open(caminho, "r", encoding="utf-8-sig", newline="") as arquivo:
        amostra = arquivo.read(4096)
        arquivo.seek(0)

        try:
            dialect = csv.Sniffer().sniff(amostra, delimiters=",;\t")
            delimitador = dialect.delimiter
        except csv.Error:
            delimitador = ","

        leitor = csv.reader(arquivo, delimiter=delimitador)
        cabecalho = next(leitor, None)

    if not cabecalho:
        raise ValueError("O CSV está vazio.")

    return {
        "table": nome_tabela_arquivo(caminho),
        "columns": criar_colunas(cabecalho),
    }


def importar_excel(caminho):
    workbook = load_workbook(
        caminho,
        read_only=True,
        data_only=True
    )

    try:
        planilha = workbook.active

        if planilha is None:
            raise ValueError("Nenhuma planilha encontrada.")

        primeira_linha = next(
            planilha.iter_rows(min_row=1, max_row=1, values_only=True),
            None
        )

        if not primeira_linha or not any(
            valor is not None for valor in primeira_linha
        ):
            raise ValueError("A primeira linha está vazia.")

        nomes = [
            valor if valor is not None else f"coluna_{i}"
            for i, valor in enumerate(primeira_linha, 1)
        ]

        return {
            "table": normalizar_nome(planilha.title),
            "columns": criar_colunas(nomes),
        }

    finally:
        workbook.close()


def importar_json(caminho):
    with open(caminho, "r", encoding="utf-8") as arquivo:
        dados = json.load(arquivo)

    # JSON de configuração do proprio programa
    if isinstance(dados, dict) and "columns" in dados:
        if not isinstance(dados["columns"], list):
            raise ValueError("A lista de colunas do JSON é inválida.")

        colunas = []

        for coluna in dados["columns"]:
            if not isinstance(coluna, dict):
                raise ValueError("Configuração de coluna inválida.")

            nome = normalizar_nome(coluna["name"])
            gerador = coluna.get("generator", "auto")

            if gerador == "idade":
                gerador = "inteiro"

                opcoes = dict(coluna.get("options") or {})
                opcoes.setdefault("min", 18)
                opcoes.setdefault("max", 90)
            else:
                opcoes = dict(coluna.get("options") or {})

            if gerador not in GERADORES:
                gerador = None

            colunas.append(
                normalizar_coluna({
                    "name": nome,
                    "type": coluna.get("type", "VARCHAR(100)"),
                    "generator": gerador,
                    "options": opcoes
                })
            )

        if not colunas:
            raise ValueError("O JSON não possui colunas.")

        return {
            "table": normalizar_nome(
                dados.get("table", nome_tabela_arquivo(caminho))
            ),
            "columns": colunas,
        }

    # JSON contendo uma lista de registros
    if isinstance(dados, list):
        if not dados or not isinstance(dados[0], dict):
            raise ValueError(
                "O JSON precisa conter uma lista de objetos."
            )

        nomes = list(dict.fromkeys(
            chave
            for registro in dados
            if isinstance(registro, dict)
            for chave in registro
        ))

        return {
            "table": normalizar_nome(nome_tabela_arquivo(caminho)),
            "columns": criar_colunas(nomes),
        }

    raise ValueError("Estrutura JSON não reconhecida.")


def listar_tabelas_sqlite(caminho):
    conexao = sqlite3.connect(
        f"file:{Path(caminho).resolve().as_posix()}?mode=ro",
        uri=True
    )

    try:
        linhas = conexao.execute("""
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
              AND name NOT LIKE 'sqlite_%'
            ORDER BY name
        """).fetchall()

        return [linha[0] for linha in linhas]

    finally:
        conexao.close()


def importar_sqlite(caminho, tabela):
    tabelas = listar_tabelas_sqlite(caminho)

    if tabela not in tabelas:
        raise ValueError("Tabela não encontrada.")

    conexao = sqlite3.connect(
        f"file:{Path(caminho).resolve().as_posix()}?mode=ro",
        uri=True
    )

    try:
        nome_seguro = tabela.replace('"', '""')

        linhas = conexao.execute(
            f'PRAGMA table_xinfo("{nome_seguro}")'
        ).fetchall()

        colunas = []

        for linha in linhas:
            nome = linha[1]
            tipo_original = (linha[2] or "TEXT").upper()
            pk = linha[5]
            hidden = linha[6]

            # ignorar colunas geradas/ocultas
            if hidden:
                continue

            # INTEGER PRIMARY KEY é preenchido automaticamente pelo SQLite
            if pk and tipo_original.strip() == "INTEGER":
                continue

            tipo_base = tipo_original.split("(")[0].strip()

            equivalencias = {
                "INTEGER": "INT",
                "SMALLINT": "INT",
                "REAL": "FLOAT",
                "NUMERIC": "DECIMAL(10,2)",
                "DECIMAL": "DECIMAL(10,2)",
                "VARCHAR": "VARCHAR(100)",
                "CHAR": "CHAR(10)",
                "BOOL": "BOOLEAN",
                "TIMESTAMP": "DATETIME",
            }

            tipo = equivalencias.get(tipo_base, tipo_base)

            if tipo not in TIPOS_SQL:
                tipo = "TEXT"

            gerador = detectar_gerador(nome)

            colunas.append({
                "name": nome,
                "type": tipo,
                "generator": gerador,
                "options": {},
            })

        if not colunas:
            raise ValueError("Nenhuma coluna utilizável encontrada.")

        return {
            "table": tabela,
            "columns": colunas,
        }

    finally:
        conexao.close()


def importar_arquivo(caminho, tabela=None):
    arquivo = Path(caminho).expanduser()

    if not arquivo.is_file():
        raise ValueError("Arquivo não encontrado.")

    extensao = arquivo.suffix.lower()

    if extensao == ".csv":
        return importar_csv(arquivo)

    if extensao == ".xlsx":
        return importar_excel(arquivo)

    if extensao == ".json":
        return importar_json(arquivo)

    if extensao == ".db":
        if not tabela:
            return {
                "tables": listar_tabelas_sqlite(arquivo)
            }

        return importar_sqlite(arquivo, tabela)

    raise ValueError(
        "Formato não suportado. Utilize JSON, DB, CSV ou XLSX."
    )
